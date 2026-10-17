"""Helper script to generate rich sample test documents (.pdf, .pptx, .docx) for NoteNest."""

import os
import docx
import fitz
import pptx
from pptx.util import Inches, Pt

DATA_DIR = os.path.join(os.path.dirname(__file__))

def create_sample_docx():
    """Create sample DOCX on AI Engineering & LLM Architecture."""
    filepath = os.path.join(DATA_DIR, "sample_ai_engineering.docx")
    doc = docx.Document()
    
    doc.add_heading("AI Engineering & Production LLM Architecture", level=0)
    
    doc.add_heading("1. Retrieval Augmented Generation (RAG) Fundamentals", level=1)
    doc.add_paragraph(
        "Retrieval Augmented Generation (RAG) optimizes the output of large language models "
        "by referencing an authoritative knowledge base outside of its training data before generating a response."
    )
    doc.add_paragraph("Key architectural stages include:", style='List Bullet')
    doc.add_paragraph("Document Chunking: Splitting source text with recursive character splitting or semantic boundary detection.", style='List Bullet')
    doc.add_paragraph("Vector Embedding: Generating dense vector embeddings using models such as text-embedding-3 or nomic-embed.", style='List Bullet')
    doc.add_paragraph("Indexing & Search: Storing vectors in HNSW / IVF indexes (e.g., Pinecone, Chroma, Milvus, pgvector).", style='List Bullet')
    doc.add_paragraph("Re-ranking & Synthesis: Employing cross-encoders (e.g. BGE-Reranker) before prompt injection.", style='List Bullet')

    doc.add_heading("2. Core Latency & Cost Trade-offs", level=1)
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Approach"
    hdr_cells[1].text = "Context Window / Latency"
    hdr_cells[2].text = "Cost & Accuracy Impact"
    
    row_data = [
        ("Naive RAG (Top-5 chunks)", "Low Latency (~300ms)", "Very Cost-Effective, Risk of missing context"),
        ("Reranked Hybrid Search", "Moderate Latency (~600ms)", "Optimal Balance, ~92% answer accuracy"),
        ("Full-Context Ingestion", "High Latency (2s - 8s)", "High Token Cost, needle-in-haystack degradation")
    ]
    for approach, latency, cost in row_data:
        row = table.add_row().cells
        row[0].text = approach
        row[1].text = latency
        row[2].text = cost

    doc.add_heading("3. Best Practices for Production Deployment", level=1)
    doc.add_paragraph(
        "Implement continuous eval harnesses (e.g. Ragas, TruLens) measuring Context Relevance, "
        "Faithfulness, and Answer Relevance. Cache recurring prompts using semantic redis caching."
    )
    
    doc.save(filepath)
    print(f"Created: {filepath}")

def create_sample_pptx():
    """Create sample PPTX on Cloud Microservices & Scalability."""
    filepath = os.path.join(DATA_DIR, "sample_cloud_architecture.pptx")
    prs = pptx.Presentation()
    blank_slide_layout = prs.slide_layouts[6]
    title_slide_layout = prs.slide_layouts[0]

    # Slide 1: Title
    slide1 = prs.slides.add_slide(title_slide_layout)
    slide1.shapes.title.text = "Modern Cloud Systems & Event-Driven Architecture"
    slide1.placeholders[1].text = "Distributed Systems, Asynchronous Messaging, and Resiliency Patterns"

    # Slide 2: Event-Driven vs Polling
    slide2 = prs.slides.add_slide(blank_slide_layout)
    title_box = slide2.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8), Inches(1))
    tf_title = title_box.text_frame
    tf_title.text = "Slide 2: Event-Driven Messaging Patterns"
    
    body_box = slide2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8), Inches(4))
    tf = body_box.text_frame
    p1 = tf.add_paragraph()
    p1.text = "Decoupling Producer and Consumer Services:"
    p1.font.bold = True
    
    p2 = tf.add_paragraph()
    p2.text = "Apache Kafka: Partitioned append-only log providing strict ordering per key and high throughput."
    p2.level = 1
    
    p3 = tf.add_paragraph()
    p3.text = "RabbitMQ: Advanced AMQP routing with exchange exchanges, fanout, and work queues."
    p3.level = 1
    
    p4 = tf.add_paragraph()
    p4.text = "Dead Letter Queues (DLQ): Safely isolate poisoned messages without halting partition consumers."
    p4.level = 1

    # Slide 3: Resiliency Strategies
    slide3 = prs.slides.add_slide(blank_slide_layout)
    title_box3 = slide3.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8), Inches(1))
    title_box3.text_frame.text = "Slide 3: High Availability & Resiliency Patterns"
    
    body_box3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8), Inches(4))
    tf3 = body_box3.text_frame
    tf3.add_paragraph().text = "- Circuit Breaker: Prevents cascading failures by short-circuiting downstream timeouts."
    tf3.add_paragraph().text = "- Idempotency Keys: Enforce at-least-once delivery semantics without duplicate side-effects."
    tf3.add_paragraph().text = "- Graceful Degradation: Fallback to cached responses or reduced functionality."

    prs.save(filepath)
    print(f"Created: {filepath}")

def create_sample_pdf():
    """Create sample PDF on Python Concurrency: Asyncio, Threading, and Multiprocessing."""
    filepath = os.path.join(DATA_DIR, "sample_python_concurrency.pdf")
    doc = fitz.open()
    
    page1 = doc.new_page()
    rect1 = fitz.Rect(50, 50, 550, 750)
    text1 = (
        "PYTHON CONCURRENCY DEEP DIVE\n\n"
        "1. THE GLOBAL INTERPRETER LOCK (GIL)\n"
        "In CPython, the Global Interpreter Lock (GIL) is a mutex that prevents multiple "
        "native threads from executing Python bytecodes at once. This lock is necessary because "
        "CPython's memory management is not thread-safe (reference counting).\n\n"
        "2. CONCURRENCY PARADIGMS COMPARED\n"
        "- Multiprocessing: Bypasses GIL by creating separate OS processes with dedicated memory heaps. "
        "Ideal for CPU-bound computations like matrix multiplication or image processing.\n"
        "- Threading: Uses OS threads sharing the same memory heap. Best suited for I/O-bound tasks "
        "like downloading web pages or querying relational databases where threads release the GIL.\n"
        "- Asyncio: Single-threaded cooperative multitasking utilizing an event loop and coroutines (async/await). "
        "Massive scalability for concurrent network sockets and WebSockets.\n\n"
        "3. PRO TIPS FOR BACKEND ENGINEERS\n"
        "Use uvloop as an event loop drop-in replacement to boost asyncio throughput by 2x-4x. "
        "Always guard against blocking calls inside async handlers by offloading to loop.run_in_executor()."
    )
    page1.insert_textbox(rect1, text1, fontsize=11, fontname="helv")

    page2 = doc.new_page()
    rect2 = fitz.Rect(50, 50, 550, 750)
    text2 = (
        "ASYNCIO INTERNALS & COMMON PITFALLS\n\n"
        "1. EVENT LOOP STARVATION\n"
        "A synchronous time.sleep() or CPU heavy loop in a coroutine freezes the entire event loop, "
        "preventing all other concurrent connections from progressing.\n\n"
        "2. TASK CANCELATION & CLEANUP\n"
        "Always handle asyncio.CancelledError when wrapping tasks in asyncio.shield() or cleanup contexts "
        "to ensure open database connections and files are properly closed.\n\n"
        "3. INTERVIEW FLASHCARD HIGHLIGHTS\n"
        "- Q: What makes asyncio lightweight compared to threads?\n"
        "- A: Coroutines do not require OS context switches or heavy thread stack allocations (typically only few hundred bytes).\n"
        "- Q: When should ProcessPoolExecutor be preferred over ThreadPoolExecutor?\n"
        "- A: Whenever the workload is CPU-bound and requires multiple CPU cores."
    )
    page2.insert_textbox(rect2, text2, fontsize=11, fontname="helv")

    doc.save(filepath)
    doc.close()
    print(f"Created: {filepath}")

if __name__ == "__main__":
    os.makedirs(DATA_DIR, exist_ok=True)
    create_sample_docx()
    create_sample_pptx()
    create_sample_pdf()
